"""Attendance: session lifecycle, the verification rules, records and export.

The clock is replaced so a test can walk through a lecture minute by minute, and
beacon evidence is built with the real token function — so each rejection test
is one real way a student could try to be marked present without being there.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.attendance import AttendanceVerification
from app.services import attendance_service, beacon
from tests.phone import Phone, ready_phone

API = "/api/v1"
START = datetime(2026, 9, 14, 9, 0, tzinfo=UTC)


# -- fixtures ------------------------------------------------------------


class Clock:
    def __init__(self) -> None:
        self.now = START

    def at(self, *, minutes: float = 0, seconds: float = 0) -> None:
        self.now = START + timedelta(minutes=minutes, seconds=seconds)


@pytest.fixture
def clock(monkeypatch: pytest.MonkeyPatch) -> Clock:
    fake = Clock()
    monkeypatch.setattr(attendance_service, "utcnow", lambda: fake.now)
    return fake


async def signup(client: AsyncClient, email: str, name: str) -> dict[str, str]:
    password = "correct-horse-battery"
    response = await client.post(
        f"{API}/auth/register",
        json={"email": email, "full_name": name, "password": password},
    )
    assert response.status_code == 201, response.text
    tokens = await client.post(f"{API}/auth/login", json={"email": email, "password": password})
    return {"Authorization": f"Bearer {tokens.json()['access_token']}"}


@pytest.fixture
async def teacher(client: AsyncClient) -> dict[str, str]:
    return await signup(client, "teacher@example.edu", "Grace Hopper")


# Each student has a verified email and a bound, signing phone (Phase 5), so
# these tests exercise the beacon rules behind a fully valid identity. The
# identity checks themselves are in test_security.py.
PHONES: dict[str, Phone] = {}


async def student_with_phone(
    client: AsyncClient, db: AsyncSession, email: str, name: str
) -> dict[str, str]:
    headers = await signup(client, email, name)
    PHONES[headers["Authorization"]] = await ready_phone(client, db, headers)
    return headers


@pytest.fixture
async def student(client: AsyncClient, db_session: AsyncSession) -> dict[str, str]:
    return await student_with_phone(client, db_session, "student@example.edu", "Alan Turing")


@pytest.fixture
async def classmate(client: AsyncClient, db_session: AsyncSession) -> dict[str, str]:
    return await student_with_phone(
        client, db_session, "classmate@example.edu", "Katherine Johnson"
    )


@pytest.fixture
async def outsider(client: AsyncClient, db_session: AsyncSession) -> dict[str, str]:
    return await student_with_phone(client, db_session, "outsider@example.edu", "Ada Lovelace")


async def join(client: AsyncClient, code: str, headers: dict[str, str]) -> None:
    joined = await client.post(f"{API}/classrooms/join", json={"code": code}, headers=headers)
    assert joined.status_code == 200, joined.text


@pytest.fixture
async def classroom(
    client: AsyncClient,
    teacher: dict[str, str],
    student: dict[str, str],
    classmate: dict[str, str],
) -> dict[str, Any]:
    created = await client.post(
        f"{API}/classrooms", json={"name": "Operating Systems"}, headers=teacher
    )
    body: dict[str, Any] = created.json()
    for headers in (student, classmate):
        await join(client, body["code"], headers)
    return body


async def start(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    **overrides: Any,
) -> dict[str, Any]:
    payload = {"date": "2026-09-14", "threshold_minutes": 5, **overrides}
    response = await client.post(
        f"{API}/classrooms/{classroom['id']}/attendance/sessions", json=payload, headers=teacher
    )
    assert response.status_code == 201, response.text
    body: dict[str, Any] = response.json()
    return body


def heard(
    session: dict[str, Any],
    windows: range | list[int],
    *,
    rssi: int = -65,
    hop: int = 0,
) -> list[dict[str, Any]]:
    """Genuine observations of the given windows."""
    session_id = uuid.UUID(session["id"])
    return [
        {
            "window": w,
            "token": beacon.token(secret=session["beacon_secret"], session_id=session_id, window=w),
            "rssi": rssi,
            "hop": hop,
        }
        for w in windows
    ]


async def verify(
    client: AsyncClient,
    session: dict[str, Any],
    headers: dict[str, str],
    observations: list[dict[str, Any]],
) -> dict[str, Any]:
    body = PHONES[headers["Authorization"]].signed_body(
        session_id=session["id"],
        observations=observations,
        issued_at=int(attendance_service.utcnow().timestamp()),
    )
    response = await client.post(
        f"{API}/attendance/sessions/{session['id']}/verify",
        json={**body, "device_id_hash": "abc123"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()
    return body


async def open_session(
    client: AsyncClient, session: dict[str, Any], teacher: dict[str, str]
) -> None:
    response = await client.patch(
        f"{API}/attendance/sessions/{session['id']}", json={"status": "active"}, headers=teacher
    )
    assert response.status_code == 200, response.text


# -- the beacon token ------------------------------------------------------


def test_token_is_pinned_so_the_mobile_codec_can_match_it() -> None:
    """The Dart implementation is tested against these exact values."""
    secret = "00" * 31 + "01"
    session_id = uuid.UUID("12345678-1234-5678-1234-567812345678")
    assert beacon.token(secret=secret, session_id=session_id, window=0) == "04eeaa24a6b04171"
    assert beacon.token(secret=secret, session_id=session_id, window=7) == "3af1fe0d4a84178e"
    assert beacon.token(secret=secret, session_id=session_id, window=123456) == "af534f2474c390fe"
    assert beacon.session_tag(session_id) == "12345678"


def test_tokens_differ_by_window_and_session() -> None:
    secret = beacon.new_secret()
    a, b = uuid.uuid4(), uuid.uuid4()
    assert beacon.token(secret=secret, session_id=a, window=1) != beacon.token(
        secret=secret, session_id=a, window=2
    )
    assert beacon.token(secret=secret, session_id=a, window=1) != beacon.token(
        secret=secret, session_id=b, window=1
    )


@pytest.mark.parametrize(
    ("elapsed", "required"), [(1, 1), (2, 2), (3, 3), (10, 7), (20, 14), (100, 70), (11, 8)]
)
def test_required_windows_is_seventy_percent_rounded_up(elapsed: int, required: int) -> None:
    assert beacon.required_windows(elapsed) == required


def test_window_index() -> None:
    assert beacon.window_index(started_at=START, at=START) == 0
    assert beacon.window_index(started_at=START, at=START + timedelta(seconds=29.9)) == 0
    assert beacon.window_index(started_at=START, at=START + timedelta(seconds=30)) == 1
    assert beacon.window_index(started_at=START, at=START - timedelta(seconds=1)) == -1


# -- session lifecycle ----------------------------------------------------


async def test_teacher_starts_a_session_with_every_student_pending(
    client: AsyncClient, classroom: dict[str, Any], teacher: dict[str, str], clock: Clock
) -> None:
    session = await start(client, classroom, teacher, hop_depth=2, rssi_threshold=-75)

    assert session["status"] == "monitoring"
    assert session["verification_opens_at"].startswith("2026-09-14T09:05:00")
    assert len(session["beacon_secret"]) == 64
    assert session["session_tag"] == uuid.UUID(session["id"]).bytes[:4].hex()
    assert session["window_seconds"] == 30
    assert (session["rssi_threshold"], session["hop_depth"]) == (-75, 2)
    assert (session["present_count"], session["record_count"]) == (0, 2)

    records = await client.get(
        f"{API}/attendance/sessions/{session['id']}/records", headers=teacher
    )
    assert {r["status"] for r in records.json()} == {"pending"}
    assert [r["student_name"] for r in records.json()] == ["Alan Turing", "Katherine Johnson"]


async def test_students_never_see_the_beacon_secret(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)

    one = await client.get(f"{API}/attendance/sessions/{session['id']}", headers=student)
    listed = await client.get(
        f"{API}/classrooms/{classroom['id']}/attendance/sessions", headers=student
    )
    assert one.json()["beacon_secret"] is None
    assert one.json()["my_status"] == "pending"
    assert [s["beacon_secret"] for s in listed.json()] == [None]


async def test_only_teachers_start_sessions(
    client: AsyncClient,
    classroom: dict[str, Any],
    student: dict[str, str],
    outsider: dict[str, str],
    clock: Clock,
) -> None:
    url = f"{API}/classrooms/{classroom['id']}/attendance/sessions"
    body = {"date": "2026-09-14"}
    assert (await client.post(url, json=body, headers=student)).status_code == 403
    assert (await client.post(url, json=body, headers=outsider)).status_code == 404


async def test_one_open_session_per_classroom(
    client: AsyncClient, classroom: dict[str, Any], teacher: dict[str, str], clock: Clock
) -> None:
    first = await start(client, classroom, teacher)
    second = await client.post(
        f"{API}/classrooms/{classroom['id']}/attendance/sessions",
        json={"date": "2026-09-14"},
        headers=teacher,
    )
    assert second.status_code == 409

    await client.patch(
        f"{API}/attendance/sessions/{first['id']}", json={"status": "ended"}, headers=teacher
    )
    await start(client, classroom, teacher)


async def test_verification_opens_by_itself_after_the_threshold(
    client: AsyncClient, classroom: dict[str, Any], teacher: dict[str, str], clock: Clock
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=5)
    url = f"{API}/attendance/sessions/{session['id']}"

    clock.at(minutes=4, seconds=59)
    assert (await client.get(url, headers=teacher)).json()["status"] == "monitoring"
    clock.at(minutes=5)
    assert (await client.get(url, headers=teacher)).json()["status"] == "active"


async def test_teacher_can_open_early_and_end(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    url = f"{API}/attendance/sessions/{session['id']}"

    assert (await client.patch(url, json={"status": "active"}, headers=student)).status_code == 403

    clock.at(minutes=1)
    opened = await client.patch(url, json={"status": "active"}, headers=teacher)
    assert opened.json()["status"] == "active"
    assert opened.json()["verification_opens_at"].startswith("2026-09-14T09:01:00")

    clock.at(minutes=50)
    ended = await client.patch(url, json={"status": "ended"}, headers=teacher)
    assert ended.json()["status"] == "ended"
    assert ended.json()["ended_at"].startswith("2026-09-14T09:50:00")

    again = await client.patch(url, json={"status": "active"}, headers=teacher)
    assert again.status_code == 409


async def test_ending_marks_pending_students_absent_including_late_joiners(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    outsider: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    # Joins after the session began, so has no pending row.
    await join(client, classroom["code"], outsider)

    await client.patch(
        f"{API}/attendance/sessions/{session['id']}", json={"status": "ended"}, headers=teacher
    )
    records = await client.get(
        f"{API}/attendance/sessions/{session['id']}/records", headers=teacher
    )
    assert len(records.json()) == 3
    assert {r["status"] for r in records.json()} == {"absent"}


# -- verification: accepted ----------------------------------------------


async def test_a_student_in_the_room_the_whole_time_is_marked_present(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    db_session: AsyncSession,
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    clock.at(minutes=6)  # window 12

    result = await verify(client, session, student, heard(session, range(13)))

    assert result["accepted"] is True
    assert result["reason"] is None
    assert (result["valid_windows"], result["elapsed_windows"]) == (13, 13)
    assert result["record"]["status"] == "present"
    assert result["record"]["marked_at"].startswith("2026-09-14T09:06:00")

    mine = await client.get(f"{API}/attendance/sessions/{session['id']}", headers=student)
    assert mine.json()["my_status"] == "present"

    row = await db_session.scalar(select(AttendanceVerification))
    assert row is not None
    assert (row.accepted, row.avg_rssi, row.hop_count, row.device_id_hash) == (
        True,
        -65,
        0,
        "abc123",
    )


async def test_exactly_seventy_percent_is_enough_and_less_is_not(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    classmate: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=4, seconds=45)  # window 9 → 10 windows begun, 7 required

    # 7 of 10, ending with a fresh reading.
    enough = await verify(client, session, student, heard(session, [0, 1, 2, 6, 7, 8, 9]))
    assert enough["accepted"] is True
    assert enough["required_windows"] == 7

    short = await verify(client, session, classmate, heard(session, [0, 1, 6, 7, 8, 9]))
    assert short["accepted"] is False
    assert short["reason"] == "insufficient_presence"
    assert (short["valid_windows"], short["required_windows"]) == (6, 7)
    assert short["record"]["status"] == "pending"


async def test_duplicate_readings_of_one_window_count_once(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=4, seconds=45)  # window 9

    padded = heard(session, [9]) * 20 + heard(session, [8], hop=1) * 20
    result = await verify(client, session, student, padded)
    assert result["reason"] == "insufficient_presence"
    assert result["valid_windows"] == 2


async def test_relayed_beacons_count_up_to_the_session_hop_depth(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    classmate: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)  # window 4

    two_hops = await verify(client, session, student, heard(session, range(5), hop=2))
    assert two_hops["accepted"] is True

    three_hops = await verify(client, session, classmate, heard(session, range(5), hop=3))
    assert three_hops["reason"] == "hop_limit_exceeded"


async def test_a_session_can_forbid_relays_altogether(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1, hop_depth=0)
    clock.at(minutes=2)
    result = await verify(client, session, student, heard(session, range(5), hop=1))
    assert result["reason"] == "hop_limit_exceeded"


async def test_marking_twice_is_idempotent(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    db_session: AsyncSession,
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    evidence = heard(session, range(5))

    first = await verify(client, session, student, evidence)
    second = await verify(client, session, student, evidence)
    assert first["record"]["id"] == second["record"]["id"]
    assert second["accepted"] is True
    assert await db_session.scalar(select(func.count()).select_from(AttendanceVerification)) == 1


async def test_a_student_who_joined_mid_session_can_still_be_marked(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    outsider: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    await join(client, classroom["code"], outsider)
    clock.at(minutes=2)

    result = await verify(client, session, outsider, heard(session, range(5)))
    assert result["accepted"] is True
    assert result["record"]["student_name"] == "Ada Lovelace"


# -- verification: every way to cheat -------------------------------------


async def test_nothing_can_be_marked_before_verification_opens(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=10)
    clock.at(minutes=9)
    result = await verify(client, session, student, heard(session, range(19)))
    assert result["accepted"] is False
    assert result["reason"] == "session_not_open"


async def test_nothing_can_be_marked_after_the_session_ends(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    await client.patch(
        f"{API}/attendance/sessions/{session['id']}", json={"status": "ended"}, headers=teacher
    )
    result = await verify(client, session, student, heard(session, range(5)))
    assert result["reason"] == "session_ended"
    assert result["record"]["status"] == "absent"


async def test_a_student_outside_the_room_hears_nothing(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    result = await verify(client, session, student, [])
    assert result["accepted"] is False
    assert result["reason"] == "no_valid_evidence"


async def test_forged_tokens_are_rejected(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    forged = [{"window": w, "token": "00" * 8, "rssi": -60, "hop": 0} for w in range(5)]
    result = await verify(client, session, student, forged)
    assert result["reason"] == "invalid_token"


async def test_tokens_from_an_earlier_session_do_not_carry_over(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    """Replaying a recording of Monday's beacon on Tuesday."""
    monday = await start(client, classroom, teacher, threshold_minutes=1)
    recorded = heard(monday, range(5))
    await client.patch(
        f"{API}/attendance/sessions/{monday['id']}", json={"status": "ended"}, headers=teacher
    )

    tuesday = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    result = await verify(client, tuesday, student, recorded)
    assert result["reason"] == "invalid_token"


async def test_tokens_for_future_windows_are_ignored(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    """A leaked secret must not let a student pre-claim the rest of the lecture."""
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)  # window 4
    evidence = heard(session, [3, 4]) + heard(session, range(5, 40))
    result = await verify(client, session, student, evidence)
    assert result["reason"] == "insufficient_presence"
    assert result["valid_windows"] == 2


async def test_weak_signals_do_not_count(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1, rssi_threshold=-80)
    clock.at(minutes=2)

    weak = await verify(client, session, student, heard(session, range(5), rssi=-81))
    assert weak["reason"] == "signal_too_weak"

    # -80 is Medium: the threshold itself counts.
    medium = await verify(client, session, student, heard(session, range(5), rssi=-80))
    assert medium["accepted"] is True


async def test_signal_must_still_be_strong_at_the_moment_of_marking(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    """In range all lecture, then walked out to the corridor before tapping Mark."""
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=5)  # window 10
    evidence = heard(session, range(10)) + heard(session, [10], rssi=-95)
    result = await verify(client, session, student, evidence)
    assert result["reason"] == "signal_too_weak"
    assert result["valid_windows"] == 10


async def test_old_evidence_alone_is_stale(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    """In range for the first half, then left — and tries to mark at the end."""
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=10)  # window 20
    result = await verify(client, session, student, heard(session, range(18)))
    assert result["reason"] == "stale_signal"

    # One window behind is still fresh: the beacon may just have rotated.
    fresh = await verify(client, session, student, heard(session, range(20)))
    assert fresh["accepted"] is True


async def test_teachers_do_not_mark_themselves(
    client: AsyncClient, classroom: dict[str, Any], teacher: dict[str, str], clock: Clock
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    response = await client.post(
        f"{API}/attendance/sessions/{session['id']}/verify",
        json={"observations": []},
        headers=teacher,
    )
    assert response.status_code == 403


async def test_malformed_evidence_is_a_validation_error(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    url = f"{API}/attendance/sessions/{session['id']}/verify"
    for bad in (
        {"window": -1, "token": "00" * 8, "rssi": -60, "hop": 0},
        {"window": 0, "token": "zz" * 8, "rssi": -60, "hop": 0},
        {"window": 0, "token": "00" * 7, "rssi": -60, "hop": 0},
    ):
        response = await client.post(url, json={"observations": [bad]}, headers=student)
        assert response.status_code == 422, bad


async def test_rejected_attempts_are_kept_with_a_reason(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher, threshold_minutes=1)
    clock.at(minutes=2)
    await verify(client, session, student, heard(session, range(5), rssi=-99))
    await verify(client, session, student, heard(session, range(5)))

    url = f"{API}/attendance/sessions/{session['id']}/verifications"
    attempts = (await client.get(url, headers=teacher)).json()
    # Sorted rather than compared in order: both rows share one transaction in
    # this test, so the database clock gives them the same `created_at`.
    assert sorted((a["accepted"], a["rejection_reason"] or "") for a in attempts) == [
        (False, "signal_too_weak"),
        (True, ""),
    ]
    assert {a["student_name"] for a in attempts} == {"Alan Turing"}
    assert (await client.get(url, headers=student)).status_code == 403


# -- records ----------------------------------------------------------------


async def test_students_see_only_their_own_record(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    records = await client.get(
        f"{API}/attendance/sessions/{session['id']}/records", headers=student
    )
    assert [r["student_name"] for r in records.json()] == ["Alan Turing"]


async def test_teachers_correct_records_by_hand(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    outsider: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    records = await client.get(
        f"{API}/attendance/sessions/{session['id']}/records", headers=teacher
    )
    record_id = records.json()[0]["id"]
    url = f"{API}/attendance/records/{record_id}"

    assert (await client.patch(url, json={"status": "present"}, headers=student)).status_code == 403
    assert (
        await client.patch(url, json={"status": "present"}, headers=outsider)
    ).status_code == 404
    assert (await client.patch(url, json={"status": "pending"}, headers=teacher)).status_code == 422

    present = await client.patch(url, json={"status": "present"}, headers=teacher)
    assert present.json()["status"] == "present"
    assert present.json()["corrected"] is True
    assert present.json()["marked_at"] is not None

    absent = await client.patch(url, json={"status": "absent"}, headers=teacher)
    assert (absent.json()["status"], absent.json()["marked_at"]) == ("absent", None)


async def test_every_attendance_route_is_404_to_a_non_member(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    outsider: dict[str, str],
    clock: Clock,
) -> None:
    session = await start(client, classroom, teacher)
    records = await client.get(
        f"{API}/attendance/sessions/{session['id']}/records", headers=teacher
    )
    sid, cid, rid = session["id"], classroom["id"], records.json()[0]["id"]

    calls = [
        ("POST", f"/classrooms/{cid}/attendance/sessions", {"date": "2026-09-14"}),
        ("GET", f"/classrooms/{cid}/attendance/sessions", None),
        ("GET", f"/attendance/sessions/{sid}", None),
        ("PATCH", f"/attendance/sessions/{sid}", {"status": "ended"}),
        ("POST", f"/attendance/sessions/{sid}/verify", {"observations": []}),
        ("GET", f"/attendance/sessions/{sid}/records", None),
        ("GET", f"/attendance/sessions/{sid}/verifications", None),
        ("PATCH", f"/attendance/records/{rid}", {"status": "present"}),
        ("GET", f"/classrooms/{cid}/attendance/export", None),
    ]
    for method, path, body in calls:
        response = await client.request(method, API + path, json=body, headers=outsider)
        assert response.status_code == 404, (method, path, response.text)


# -- export -------------------------------------------------------------------


async def test_export_matches_the_old_spreadsheet(
    client: AsyncClient,
    classroom: dict[str, Any],
    teacher: dict[str, str],
    student: dict[str, str],
    clock: Clock,
) -> None:
    # Two ended sessions: Alan present at one, absent at the other.
    for present in (True, False):
        session = await start(client, classroom, teacher, threshold_minutes=1)
        clock.at(minutes=2)
        if present:
            await verify(client, session, student, heard(session, range(5)))
        await client.patch(
            f"{API}/attendance/sessions/{session['id']}", json={"status": "ended"}, headers=teacher
        )
        clock.at()
    # An open session must not count against anyone.
    await start(client, classroom, teacher)

    url = f"{API}/classrooms/{classroom['id']}/attendance/export"
    assert (await client.get(url, headers=student)).status_code == 403

    response = await client.get(url, headers=teacher)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        "Operating%20Systems_Attendance_2026-09-14.xlsx" in response.headers["content-disposition"]
    )

    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    assert sheet.title == "Attendance"
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    assert rows == [
        ["Student Name", "Email", "Present", "Absent", "Total Sessions", "Attendance %"],
        ["Alan Turing", "student@example.edu", 1, 1, 2, "50.0%"],
        ["Katherine Johnson", "classmate@example.edu", 0, 2, 2, "0.0%"],
    ]


async def test_export_neutralises_formula_injection(client: AsyncClient) -> None:
    body = attendance_service.build_xlsx([['=HYPERLINK("http://evil")', "a@b.c", 0, 0, 0, "0%"]])
    sheet = load_workbook(io.BytesIO(body)).active
    assert sheet is not None
    assert sheet["A2"].value == '\'=HYPERLINK("http://evil")'
    assert sheet["A2"].data_type == "s"
